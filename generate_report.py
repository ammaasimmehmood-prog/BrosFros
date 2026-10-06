import docx
import os
import re
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def process_bold(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)

template_path = r"C:\Users\Aasim Mehmood Mirza\OneDrive\Desktop\BrosFros (Chromium, Firefox, IE) v2\Project Report\GHUFRAN_REPORT (1).docx"
out_path = r"C:\Users\Aasim Mehmood Mirza\OneDrive\Desktop\BrosFros (Chromium, Firefox, IE) v2\Project Report\Aasim M Mirza Report.docx"
md_path = r"C:\Users\Aasim Mehmood Mirza\OneDrive\Desktop\BrosFros (Chromium, Firefox, IE) v2\Detailed_MTech_Project_Report.md"

doc = docx.Document(template_path)

# Clear existing content but keep section properties (margins, headers, footers)
body = doc._element.body
for element in list(body):
    if not element.tag.endswith('sectPr'):
        body.remove(element)

with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

image_paths = [
    r"C:\Users\Aasim Mehmood Mirza\.gemini\antigravity\brain\73cf14d4-462a-4186-ae2e-36bf033c934e\media__1783239354620.png",
    r"C:\Users\Aasim Mehmood Mirza\.gemini\antigravity\brain\73cf14d4-462a-4186-ae2e-36bf033c934e\media__1783239354721.png",
    r"C:\Users\Aasim Mehmood Mirza\.gemini\antigravity\brain\73cf14d4-462a-4186-ae2e-36bf033c934e\media__1783239354727.png",
    r"C:\Users\Aasim Mehmood Mirza\.gemini\antigravity\brain\73cf14d4-462a-4186-ae2e-36bf033c934e\media__1783239354730.png",
    r"C:\Users\Aasim Mehmood Mirza\.gemini\antigravity\brain\73cf14d4-462a-4186-ae2e-36bf033c934e\media__1783239354739.png"
]

images_added = False

for line in lines:
    line = line.strip()
    if line == "<div style=\"page-break-after: always;\"></div>":
        doc.add_page_break()
        continue
    
    if line.startswith("### "):
        doc.add_heading(line[4:], level=3)
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=2)
    elif line.startswith("# "):
        doc.add_heading(line[2:], level=1)
    elif line.startswith("- "):
        p = doc.add_paragraph(style='List Paragraph')
        process_bold(p, line)
    else:
        p = doc.add_paragraph()
        process_bold(p, line)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
    # Insert screenshots in Chapter 7
    if "CHAPTER 7: TESTING & QUALITY ASSURANCE" in line:
        doc.add_heading("7.4 System Working & Testing (Screenshots)", level=2)
        p_intro = doc.add_paragraph()
        p_intro.add_run("The following screenshots demonstrate the working and testing phases of the BrosFros Forensic Suite.").bold = False
        
        for idx, img_path in enumerate(image_paths):
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p_img.add_run()
                run.add_picture(img_path, width=Inches(6.0))
                
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_run = p_cap.add_run(f"Figure 7.{idx+1}: System Interface and Testing")
                cap_run.italic = True
        images_added = True

# Add References if they are not correctly formatted or something
# The markdown already has references.

doc.save(out_path)
print("Saved to", out_path)
